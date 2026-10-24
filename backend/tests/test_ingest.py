import io
import json
from datetime import datetime
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from pypdf import PdfWriter

from app.services.ingest import parse_upload


ROOT = Path(__file__).resolve().parents[2]


def _pdf_with_text(text: str) -> bytes:
    stream = f"BT /F1 18 Tf 72 100 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Count 1 /Kids [3 0 R] >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 300 200] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{index} 0 obj\n".encode() + body + b"\nendobj\n"
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n".encode()
    out += b"0000000000 65535 f \n"
    for offset in offsets[1:]:
        out += f"{offset:010d} 00000 n \n".encode()
    out += f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return bytes(out)


def _blank_pdf() -> bytes:
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


def _docx_bytes() -> bytes:
    document = Document()
    document.add_paragraph("说明文档，没有订单变更。")
    grid = document.add_table(rows=2, cols=2)
    grid.rows[0].cells[0].text = "物料"
    grid.rows[0].cells[1].text = "数量"
    grid.rows[1].cells[0].text = "密封圈"
    grid.rows[1].cells[1].text = "3"
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _xlsx_bytes() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.append(["物料", "数量"])
    sheet.append(["密封圈", 3])
    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def test_xlsx_keeps_dates_and_skips_a_title_row():
    book = Workbook()
    sheet = book.active
    sheet.append(["南浦配电订货表"])
    sheet.append(["物料", "规格", "数量", "订单交期"])
    sheet.append(["塑壳断路器", "250A 3P", 40.0, datetime(2026, 11, 20)])
    buffer = io.BytesIO()
    book.save(buffer)
    text, table, kind = parse_upload("订货表.xlsx", buffer.getvalue())
    assert kind == "table"
    assert table["columns"][:4] == ["物料", "规格", "数量", "订单交期"]
    assert table["rows"] == [["塑壳断路器", "250A 3P", "40", "2026-11-20"]]
    assert "2026-11-20 00:00:00" not in text
    assert "南浦配电订货表" in text


def test_parser_routes_local_files():
    text, table, kind = parse_upload("库存.csv", "物料,数量\n密封圈,3\n".encode("gb18030"))
    assert kind == "table"
    assert table["rows"][0] == ["密封圈", "3"]
    assert "[R2C1] 密封圈" in text

    text, table, kind = parse_upload("库存.xlsx", _xlsx_bytes())
    assert kind == "table"
    assert table["columns"] == ["物料", "数量"]
    assert "3" in text

    text, table, kind = parse_upload("说明.docx", _docx_bytes())
    assert kind == "table"
    assert "说明文档" in text
    assert "[R2C2] 3" in text

    text, table, kind = parse_upload("说明.pdf", _pdf_with_text("Stock120"))
    assert kind == "document"
    assert "Stock120" in text
    assert table == {}

    chat = "[2026-11-12 09:10] 小周\n50件已经点完。\n".encode()
    text, _, kind = parse_upload("记录.txt", chat)
    assert kind == "chat"
    assert "小周" in text

    payload = json.dumps(
        {"messages": [{"sender": "小周", "time": "2026-11-12 09:10", "text": "50件已经点完"}]},
        ensure_ascii=False,
    ).encode()
    text, _, kind = parse_upload("wechat.json", payload)
    assert kind == "chat"
    assert "小周" in text
    assert "50件已经点完" in text

    html = "<html><body><h1>微信聊天记录</h1><div>小周：50件已经点完</div><script>secret()</script></body></html>"
    text, _, kind = parse_upload("导出.html", html.encode())
    assert kind == "chat"
    assert "小周" in text
    assert "secret" not in text


def test_plain_text_file_stays_untyped_until_the_caller_decides():
    text, table, kind = parse_upload("说明.txt", "这是一份说明，没有订单变更。".encode())
    assert kind is None
    assert "说明" in text
    assert table == {}


def test_image_and_audio_and_scanned_pdf_fail_without_keys(monkeypatch):
    monkeypatch.setenv("OCR_API_KEY", "")
    monkeypatch.setenv("ASR_API_KEY", "")

    def boom(*args, **kwargs):
        raise AssertionError("不应访问外部接口")

    monkeypatch.setattr("app.services.media.httpx.post", boom)
    with pytest.raises(ValueError, match="OCR_API_KEY"):
        parse_upload("scan.png", b"not-an-image")
    with pytest.raises(ValueError, match="OCR_API_KEY"):
        parse_upload("scan.pdf", _blank_pdf())
    with pytest.raises(ValueError, match="ASR_API_KEY"):
        parse_upload("call.wav", b"not-audio")


def test_unsupported_extension_names_the_local_types():
    with pytest.raises(ValueError, match="XLSX"):
        parse_upload("归档.zip", b"PK")
    with pytest.raises(ValueError, match="XLSX"):
        parse_upload("旧表.xls", b"not-xlsx")


def test_review_import_filter_is_chat_or_file():
    review = (ROOT / "frontend" / "src" / "views" / "Review.vue").read_text(encoding="utf-8")
    labels = (ROOT / "frontend" / "src" / "utils" / "labels.ts").read_text(encoding="utf-8")
    assert 'label: "聊天记录"' in labels
    assert 'label: "文件"' in labels
    assert "importKindOptions" in review
    assert "备注" not in review
    assert 'note: "备注"' in labels


def test_upload_docx_and_file_text_and_reject_image_without_key(client):
    order = client.get("/orders").json()
    target = next(item for item in order if item["order_no"] == "ORD-018")
    docx = client.post(
        f"/orders/{target['id']}/sources",
        data={"type": "file", "title": "说明"},
        files={"file": ("说明.docx", _docx_bytes(), "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
    )
    assert docx.status_code == 200, docx.text
    evidence = client.get(f"/evidence/{docx.json()['source_id']}").json()
    assert evidence["type"] == "table"
    assert "[R2C1] 密封圈" in evidence["raw_text"]

    plain = client.post(
        f"/orders/{target['id']}/sources",
        data={"type": "file", "title": "纯文本"},
        files={"file": ("说明.txt", "这是一份说明，没有订单变更。".encode(), "text/plain")},
    )
    assert plain.status_code == 200, plain.text
    assert client.get(f"/evidence/{plain.json()['source_id']}").json()["type"] == "document"

    image = client.post(
        f"/orders/{target['id']}/sources",
        data={"type": "file", "title": "图片"},
        files={"file": ("scan.png", b"not-an-image", "image/png")},
    )
    assert image.status_code == 400
    assert "OCR_API_KEY" in image.json()["detail"]
    assert "source_id" not in image.json()
