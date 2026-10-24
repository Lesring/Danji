import csv
import io
import json
import re
from datetime import date, datetime
from html.parser import HTMLParser
from pathlib import Path

import openpyxl
from docx import Document
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.config import get_settings
from app.services.media import ocr_image, transcribe_audio

SOURCE_TYPES = {"chat", "table", "call_transcript", "note", "document"}

IMAGE_TYPES = {
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".webp": "image/webp",
    ".bmp": "image/bmp",
    ".gif": "image/gif",
}
AUDIO_TYPES = {".mp3", ".wav", ".m4a", ".ogg", ".aac", ".webm"}
UNSUPPORTED = "只支持 TXT、HTML、JSON、CSV、XLSX、DOCX、PDF、PNG、JPG、WEBP、MP3、WAV、M4A"

_BRACKET_CHAT = re.compile(r"^\[\d{4}[-/.]\d{1,2}[-/.]\d{1,2}[ T]\d{1,2}:\d{2}")
_WECHAT_CHAT = re.compile(r"^\d{4}-\d{2}-\d{2}\s+\d{2}:\d{2}:\d{2}\s+\S+")


_HEADER_ALIASES = (
    "物料",
    "商品",
    "产品",
    "品名",
    "订单数量",
    "数量",
    "件数",
    "规格",
    "型号",
    "订单交期",
    "交货日期",
    "交期",
    "含税单价",
    "单价",
    "订单号",
    "在手数量",
    "最早可完工日期",
)


def table_to_text(columns: list[str], rows: list[list[str]]) -> str:
    lines = [" | ".join(f"[R1C{index + 1}] {cell}" for index, cell in enumerate(columns))]
    for row_index, row in enumerate(rows, start=2):
        lines.append(
            " | ".join(f"[R{row_index}C{col_index + 1}] {cell}" for col_index, cell in enumerate(row))
        )
    return "\n".join(lines)


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.hour == value.minute == value.second == 0:
            return value.strftime("%Y-%m-%d")
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, bool):
        return "是" if value else "否"
    return str(value).strip()


def _header_index(rows: list[list[str]]) -> int:
    for index, row in enumerate(rows[:8]):
        compact = [re.sub(r"\s+", "", cell) for cell in row]
        hits = sum(1 for cell in compact if any(alias in cell for alias in _HEADER_ALIASES))
        if hits >= 2:
            return index
    return 0


def _split_table(rows: list[list[str]]) -> tuple[list[str], list[str], list[list[str]]]:
    if not rows:
        return [], [], []
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    header_at = _header_index(padded)
    titles = [" ".join(cell for cell in row if cell) for row in padded[:header_at]]
    columns = padded[header_at]
    body = [row for row in padded[header_at + 1 :] if any(cell.strip() for cell in row)]
    return titles, columns, body


def parse_csv(data: bytes) -> tuple[str, dict]:
    text = _decode_text(data)
    reader = csv.reader(io.StringIO(text))
    rows = [[cell.strip() for cell in row] for row in reader if any(cell.strip() for cell in row)]
    return _table_payload(rows)


def parse_xlsx(data: bytes) -> tuple[str, dict]:
    workbook = openpyxl.load_workbook(io.BytesIO(data), data_only=True)
    sheets = []
    parts: list[str] = []
    for sheet in workbook.worksheets:
        rows: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = [_cell_text(cell) for cell in row]
            if any(cells):
                rows.append(cells)
        titles, columns, body = _split_table(rows)
        if not columns:
            continue
        sheets.append({"sheet": sheet.title, "columns": columns, "rows": body})
        block = []
        if len(workbook.worksheets) > 1:
            block.append(sheet.title)
        block.extend(titles)
        block.append(table_to_text(columns, body))
        parts.append("\n".join(line for line in block if line))
    if not sheets:
        return "", {"columns": [], "rows": [], "sheets": []}
    return "\n\n".join(parts), {"columns": sheets[0]["columns"], "rows": sheets[0]["rows"], "sheets": sheets}


def _table_payload(rows: list[list[str]]) -> tuple[str, dict]:
    titles, columns, body = _split_table(rows)
    if not columns:
        return "", {"columns": [], "rows": []}
    text = "\n".join([*titles, table_to_text(columns, body)]).strip()
    return text, {"columns": columns, "rows": body}


def parse_upload(filename: str, data: bytes) -> tuple[str, dict, str | None]:
    if not data:
        raise ValueError("文件是空的")
    ext = Path(filename or "").suffix.lower()
    if ext == ".csv":
        return _done(*parse_csv(data), "table")
    if ext in {".xlsx", ".xlsm"}:
        try:
            text, table = parse_xlsx(data)
        except Exception as exc:
            raise ValueError("XLSX 无法读取，请确认文件未损坏。") from exc
        return _done(text, table, "table")
    if ext == ".xls":
        raise ValueError("暂不支持旧版 XLS，请另存为 XLSX 后再上传。")
    if ext == ".docx":
        return _done(*parse_docx(data))
    if ext == ".doc":
        raise ValueError("暂不支持旧版 DOC，请另存为 DOCX 后再上传。")
    if ext == ".pdf":
        return _done(*parse_pdf(data))
    if ext in IMAGE_TYPES:
        return _done(ocr_image(data, IMAGE_TYPES[ext]), {}, "document")
    if ext in AUDIO_TYPES:
        return _done(transcribe_audio(data, Path(filename).name), {}, "call_transcript")
    if ext in {".txt", ".md", ".log", ".text"}:
        text = _decode_text(data)
        kind = "chat" if looks_like_chat(filename, text) else None
        return _done(text, {}, kind)
    if ext in {".html", ".htm"}:
        text = html_to_text(_decode_text(data))
        kind = "chat" if looks_like_chat(filename, text) else "document"
        return _done(text, {}, kind)
    if ext == ".json":
        return _done(*parse_chat_json(data))
    raise ValueError(UNSUPPORTED)


def looks_like_chat(filename: str, text: str) -> bool:
    name = (filename or "").lower()
    if any(token in name for token in ("chat", "wechat", "weixin")) or any(token in (filename or "") for token in ("聊天", "微信")):
        return True
    if "微信聊天记录" in text or "QQ聊天记录" in text:
        return True
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    if any(_BRACKET_CHAT.match(line) for line in lines):
        return True
    return sum(1 for line in lines if _WECHAT_CHAT.match(line)) >= 2


def parse_docx(data: bytes) -> tuple[str, dict, str]:
    try:
        document = Document(io.BytesIO(data))
    except (PackageNotFoundError, ValueError) as exc:
        raise ValueError("DOCX 无法读取，请确认文件未损坏。") from exc
    blocks = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    table: dict = {}
    for grid in document.tables:
        rows = _table_rows(grid)
        if not rows:
            continue
        width = max(len(row) for row in rows)
        rows = [row + [""] * (width - len(row)) for row in rows]
        columns, body = rows[0], rows[1:]
        blocks.append(table_to_text(columns, body))
        if not table:
            table = {"columns": columns, "rows": body}
    kind = "table" if table else "document"
    return "\n".join(blocks), table, kind


def parse_pdf(data: bytes) -> tuple[str, dict, str]:
    try:
        reader = PdfReader(io.BytesIO(data))
    except PdfReadError as exc:
        raise ValueError("PDF 无法读取，可能已加密或已损坏。") from exc
    pages: list[str] = []
    for index, page in enumerate(reader.pages, start=1):
        extracted = page.extract_text() or ""
        if extracted.strip():
            pages.append(f"[第{index}页]\n{extracted.strip()}")
    text = "\n\n".join(pages).strip()
    if text:
        return text, {}, "document"
    settings = get_settings()
    if settings.ocr_api_key:
        try:
            from app.services.media import read_file

            ark_text = read_file(
                data,
                "document.pdf",
                "提取文档中的全部文字，保留换行和表格结构。不要解释。",
            )
            if str(ark_text).strip():
                return str(ark_text).strip(), {}, "document"
        except ValueError:
            pass
    images = _pdf_images(reader)
    if images and settings.ocr_api_key:
        parts = [ocr_image(blob, mime) for blob, mime in images]
        return "\n\n".join(part for part in parts if part.strip()), {}, "document"
    if not get_settings().ocr_api_key:
        raise ValueError(
            "这份 PDF 没有可复制的文字。识别扫描件需要在 backend/.env 填写 OCR_API_KEY。"
            "带文字层的 PDF、TXT、CSV、XLSX、DOCX 不需要这个密钥。"
        )
    raise ValueError("这份 PDF 没有可复制的文字，也没有可送给识别接口的内嵌图片。请把扫描页另存为 PNG 或 JPG。")


def parse_chat_json(data: bytes) -> tuple[str, dict, str]:
    try:
        payload = json.loads(_decode_text(data))
    except json.JSONDecodeError as exc:
        raise ValueError("JSON 无法解析。") from exc
    messages = _chat_messages(payload)
    if not messages:
        raise ValueError("这个 JSON 不是可识别的聊天记录。需要包含消息数组，每条有正文。")
    lines: list[str] = []
    for item in messages:
        head = " ".join(part for part in (f"[{item['time']}]" if item["time"] else "", item["sender"]) if part).strip()
        lines.append(f"{head}\n{item['text']}" if head else item["text"])
    return "\n\n".join(lines), {}, "chat"


def locate_cell(text: str, excerpt: str) -> tuple[int | None, int | None]:
    if not excerpt:
        return None, None
    for line in text.splitlines():
        index = line.find(excerpt)
        if index < 0:
            continue
        chosen = None
        for marker in re.finditer(r"\[R(\d+)C(\d+)\]", line):
            if marker.start() <= index:
                chosen = marker
        if chosen:
            return int(chosen.group(1)), int(chosen.group(2))
    return None, None


def _done(text: str, table: dict, kind: str | None) -> tuple[str, dict, str | None]:
    if not str(text).strip():
        raise ValueError("文件里没有可读内容")
    return text, table, kind


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("文件不是 UTF-8 或 GB18030 文本")


def _table_rows(grid) -> list[list[str]]:
    rows: list[list[str]] = []
    for row in grid.rows:
        cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
        if any(cells):
            rows.append(cells)
    return rows


def _pdf_images(reader: PdfReader) -> list[tuple[bytes, str]]:
    found: list[tuple[bytes, str]] = []
    for page in reader.pages:
        for image in getattr(page, "images", []) or []:
            blob = getattr(image, "data", None)
            if not blob:
                continue
            name = str(getattr(image, "name", "") or "").lower()
            mime = "image/png" if name.endswith(".png") else "image/jpeg"
            found.append((blob, mime))
    return found


class _HTMLText(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self._skip = 0
        self.parts: list[str] = []

    def handle_starttag(self, tag: str, attrs) -> None:
        if tag in {"script", "style"}:
            self._skip += 1

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"} and self._skip:
            self._skip -= 1
        if tag in {"p", "div", "br", "tr", "li", "h1", "h2", "h3"}:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._skip or not data.strip():
            return
        self.parts.append(data.strip())
        self.parts.append("\n")


def html_to_text(raw: str) -> str:
    parser = _HTMLText()
    parser.feed(raw)
    return "".join(parser.parts).strip()


def _chat_messages(payload) -> list[dict]:
    rows = payload
    if isinstance(payload, dict):
        rows = None
        for key in ("messages", "chat", "msgs", "data"):
            if isinstance(payload.get(key), list):
                rows = payload[key]
                break
        if rows is None:
            return []
    if not isinstance(rows, list):
        return []
    found: list[dict] = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        text = ""
        for key in ("text", "content", "message", "msg"):
            if row.get(key):
                text = str(row[key]).strip()
                break
        if not text:
            continue
        sender = ""
        for key in ("sender", "name", "nickname", "from", "user"):
            if row.get(key):
                sender = str(row[key]).strip()
                break
        stamp = ""
        for key in ("time", "timestamp", "datetime", "date"):
            if row.get(key):
                stamp = str(row[key]).strip()
                break
        found.append({"text": text, "sender": sender, "time": stamp})
    return found
